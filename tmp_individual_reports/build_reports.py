from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Asgm_SDS_V2")
OUT = ROOT / "output" / "individual_reports"
OUT.mkdir(parents=True, exist_ok=True)

MODELS = {
    "ets": {
        "student": "Chan YH",
        "model": "ETS(A,N,A)",
        "script": "03_ChanYH_ets.R",
        "file": "RDS2S1G6_Individual_Report_ChanYH_ETS.docx",
        "fig1": ROOT / "output/plots/group_summary/ets_states.png",
        "fig1_caption": "Figure 1. ETS(A,N,A) estimated level and seasonal states",
        "fig2": ROOT / "output/plots/group_summary/fc_ets.png",
        "fig2_caption": "Figure 2. ETS twelve month test window forecast against actual rainfall",
        "fig3": ROOT / "output/plots/group_summary/resid_ets.png",
        "fig3_caption": "Figure 3. Residual diagnostics for the ETS model",
        "method": [
            "This report analyses monthly mean precipitation in Kuala Lumpur, measured in millimetres per day, using NASA POWER data from January 1981 to December 2025. The series contains 540 monthly observations and no missing values after the data-quality check. The exploratory decomposition indicates moderate annual seasonality and only a weak long-run trend. The raw-series Ljung-Box test at lag 24 is strongly significant, so a useful forecasting model must represent the recurring twelve-month pattern and the remaining serial dependence rather than treat rainfall as independent observations.",
            "I selected ETS(A,N,A), an exponential-smoothing model with additive error, no trend and additive monthly seasonality. This form is appropriate because the decomposition does not support a persistent trend, while the seasonal pattern is regular enough to be represented by twelve additive seasonal states. The fitted level smoothing parameter was 0.017 and the seasonal smoothing parameter was approximately 0.0001. These small values mean that the model updates cautiously and lets the established monthly pattern carry most of the forecast. The final 12 months were reserved as a time-ordered test set; the remaining 528 months were used for estimation. Rolling-origin cross-validation then used 27 forecast origins, keeping the test year separate from model selection."
        ],
        "analysis": [
            "Performance was evaluated using RMSE, MAE and MASE, with seasonal naive forecasting as the benchmark. A MASE below one indicates that the model improves on repeating the rainfall value from the same month in the previous year. ETS obtained a training RMSE of 2.436 and a training MASE of 0.738. In rolling-origin validation, RMSE was 2.566 and MASE was 0.811, compared with 0.985 for seasonal naive forecasting. Thus, the ETS model reduced cross-validated scaled error by about 17.7 percent while its RMSE increased only 5.3 percent relative to the training estimate. The held-out test MAE was 2.077 mm/day and test MASE was 0.807, which remains below the seasonal-naive test MASE of 0.942.",
            "The residual evidence is consistent with an adequate monthly forecasting model. The lag-24 Ljung-Box p-value is 0.122, so the residuals do not show statistically significant autocorrelation at the 5 percent level. Figure 3 shows residuals centred around zero and a distribution that is broadly bell-shaped. Only one of the first 24 autocorrelations is outside its approximate confidence band. Figure 2 shows that the forecast follows the usual seasonal movement, although unusually high or low monthly rainfall remains difficult to reproduce exactly."
        ],
        "conclusion": [
            "ETS(A,N,A) provides a compact and interpretable forecast for monthly Kuala Lumpur rainfall. It captures the annual pattern, outperforms the seasonal-naive benchmark in both validation and the held-out year, and leaves no significant residual autocorrelation at lag 24. The model is therefore suitable for routine monthly planning where an estimate of the expected rainfall level is required.",
            "Its limitations are equally important. Additive seasonality smooths extreme wet months, so the model should not be interpreted as an extreme-rainfall or flood-warning system. In addition, the data represent one gridded location and include no atmospheric covariates. Future work could compare damped-trend and multiplicative-error alternatives, incorporate climate indicators such as ENSO, and test the model at several locations across Greater Kuala Lumpur."
        ],
        "metrics": [("MASE", "0.738", "0.811", "9.9%"), ("RMSE", "2.436", "2.566", "1.053"), ("Ljung Box p (lag-24)", "0.122", "-", "-")],
        "refs": [
            "Hyndman, R. J., Koehler, A. B., Ord, J. K., & Snyder, R. D. (2008). Forecasting with exponential smoothing: The state space approach. Springer.",
            "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
            "NASA Langley Research Center. (n.d.). NASA POWER: Prediction of worldwide energy resources. https://power.larc.nasa.gov/",
            "United Nations. (n.d.). Sustainable Development Goal 13: Climate action. https://sdgs.un.org/goals/goal13"
        ]
    },
    "arima": {
        "student": "Steph QF",
        "model": "ARIMA + Fourier", "script": "04_StephQF_arima.R",
        "file": "RDS2S1G6_Individual_Report_StephQF_ARIMA_Fourier.docx",
        "fig1": ROOT / "output/plots/group_summary/arima_differenced_acf_pacf.png",
        "fig1_caption": "Figure 1. ACF and PACF after differencing the rainfall series",
        "fig2": ROOT / "output/plots/group_summary/fc_arima.png",
        "fig2_caption": "Figure 2. ARIMA plus Fourier twelve month test window forecast against actual rainfall",
        "fig3": ROOT / "output/plots/group_summary/resid_arima.png",
        "fig3_caption": "Figure 3. Residual diagnostics for the ARIMA plus Fourier model",
        "method": [
            "The analysis uses the NASA POWER monthly mean precipitation rate for Kuala Lumpur from January 1981 to December 2025. There are 540 observations with no missing values after checking and interpolation. The series displays a repeated annual cycle, a weak trend and considerable short-term variability. A lag-24 Ljung-Box test on the unmodelled series rejects white noise, confirming that a forecast should account for both annual seasonality and temporal correlation. The last 12 observations were held out for final testing, leaving 528 months for model estimation.",
            "I selected a dynamic regression model with four Fourier pairs for the annual cycle and ARIMA(0,1,1) errors. Fourier terms provide a parsimonious smooth representation of seasonality: four sine and cosine pairs replace a full set of separate monthly indicators. Differencing addresses the non-stationary level, while the moving-average error term captures the short-run dependence left after seasonal adjustment. The estimated MA(1) coefficient was about -0.977, indicating that consecutive forecast errors are strongly offset. Model evaluation used 27 rolling-origin validation folds in addition to the fixed 2025 test set, so the selected specification was assessed on observations not used in fitting."
        ],
        "analysis": [
            "The ARIMA plus Fourier model achieved a training RMSE of 2.438 and MASE of 0.738. Its rolling-origin RMSE was 2.538 and MASE was 0.801, against 3.179 and 0.985 respectively for the seasonal-naive benchmark. The cross-validated MASE is therefore about 18.7 percent lower than the benchmark, and the RMSE ratio of 1.041 suggests little loss when the model is applied to new data. The final-year test RMSE was 2.424, test MAE was 2.170 mm/day and test MASE was 0.843. Although the test year includes irregular monthly variation, the model still improves on seasonal naive forecasting, whose test MASE was 0.942.",
            "Residual diagnostics support the fitted specification. The lag-24 Ljung-Box p-value is 0.108, above the 5 percent threshold, and only one of the first 24 ACF lags lies outside the confidence limits. The residual plot fluctuates around zero without a sustained sequence, while the histogram is approximately symmetric with some upper-tail rainfall errors. Figure 2 shows that the fitted seasonal shape tracks ordinary months well; the largest peaks remain more difficult because a smooth Fourier representation deliberately avoids chasing isolated shocks."
        ],
        "conclusion": [
            "The ARIMA plus Fourier model combines a transparent seasonal structure with an error process that handles remaining short-run correlation. It performs better than seasonal naive forecasting in cross-validation and in the held-out year, and its residual diagnostics do not indicate significant autocorrelation at lag 24. This makes it a credible model for monthly rainfall planning and comparison against more flexible approaches.",
            "The model still has limitations. A fixed set of Fourier terms assumes that the average seasonal shape changes smoothly, so rare convective or monsoon extremes may be understated. The differencing step also makes long-run interpretation less direct. Further work could add weather or climate predictors, assess time-varying Fourier coefficients, and evaluate interval calibration for high-impact rainfall months."
        ],
        "metrics": [("MASE", "0.738", "0.801", "8.5%"), ("RMSE", "2.438", "2.538", "1.041"), ("Ljung Box p (lag-24)", "0.108", "-", "-")],
        "refs": [
            "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
            "Hyndman, R. J., & Khandakar, Y. (2008). Automatic time series forecasting: The forecast package for R. Journal of Statistical Software, 27(3), 1-22. https://doi.org/10.18637/jss.v027.i03",
            "NASA Langley Research Center. (n.d.). NASA POWER: Prediction of worldwide energy resources. https://power.larc.nasa.gov/",
            "United Nations. (n.d.). Sustainable Development Goal 13: Climate action. https://sdgs.un.org/goals/goal13"
        ]
    },
    "tslm": {
        "student": "Ham GQ", "model": "TSLM + Fourier", "script": "05_HamGQ_tslm.R",
        "file": "RDS2S1G6_Individual_Report_HamGQ_TSLM_Fourier.docx",
        "fig1": ROOT / "output/plots/group_summary/tslm_fit_diagnostics.png",
        "fig1_caption": "Figure 1. TSLM fitted-value diagnostics for the Kuala Lumpur rainfall series",
        "fig2": ROOT / "output/plots/group_summary/fc_tslm.png",
        "fig2_caption": "Figure 2. TSLM plus Fourier twelve month test window forecast against actual rainfall",
        "fig3": ROOT / "output/plots/group_summary/resid_tslm.png",
        "fig3_caption": "Figure 3. Residual diagnostics for the TSLM plus Fourier model",
        "method": [
            "This report forecasts monthly mean precipitation in Kuala Lumpur using NASA POWER observations from January 1981 through December 2025. The cleaned data contain 540 consecutive monthly values. Exploratory analysis identifies an annual cycle as the most visible feature, while the long-run trend is much weaker. The raw series is autocorrelated at seasonal lags and rainfall occasionally shows large positive departures. These characteristics support a regression model that can estimate both a gradual time effect and a smooth seasonal pattern. The final 12 months were set aside for testing and the preceding 528 months formed the training sample.",
            "I used a time-series linear model (TSLM) with a linear trend and Fourier terms with K = 4. The specification estimates one linear time coefficient and eight trigonometric regressors, four sine and four cosine terms, for the annual period of 12 months. The trend estimate was 0.00348 mm/day per month, which is small but positive. This model was chosen because it is interpretable: the trend shows the average long-run movement and the Fourier terms show the recurring annual pattern without creating twelve separate dummy coefficients. Twenty-seven rolling-origin folds were used to evaluate how consistently the regression forecasts future months."
        ],
        "analysis": [
            "The TSLM plus Fourier model produced a training RMSE of 2.417 and MASE of 0.736. Under rolling-origin validation, RMSE was 2.557 and MASE was 0.814. The seasonal-naive benchmark produced a cross-validated MASE of 0.985, so the regression lowers scaled error by about 17.4 percent. The ratio of cross-validated to training RMSE is 1.058, which indicates a modest and acceptable decline outside the fitting sample. In the 2025 test period, RMSE was 2.407, MAE was 2.146 mm/day and MASE was 0.833, again below the seasonal-naive test MASE of 0.942.",
            "The error checks are encouraging. The lag-24 Ljung-Box p-value is 0.148 and only one of the first 24 ACF bars exceeds its approximate confidence bound. The residuals are centred near zero, with a roughly bell-shaped central distribution and a small number of high-rainfall outliers. Figure 2 indicates that the model represents normal seasonal changes well, but a linear trend plus fixed Fourier waves cannot reproduce every abrupt monthly peak. This is an expected trade-off for a model designed for clarity and stable average forecasts."
        ],
        "conclusion": [
            "The TSLM plus Fourier specification is a useful interpretable baseline for monthly rainfall forecasting. It improves on seasonal naive forecasting in rolling validation and the held-out year, while its residuals pass the lag-24 whiteness check. The explicit trend and seasonal coefficients also make the model easy to communicate to non-technical planning users.",
            "However, the linear trend and fixed Fourier cycle cannot react to sudden weather regimes or rare extremes. The model also excludes explanatory information such as sea-surface temperature, humidity and regional circulation. Future work should test dynamic regression with climate covariates, robust or transformed responses for extreme months, and forecasts across multiple stations before using the results for regional decisions."
        ],
        "metrics": [("MASE", "0.736", "0.814", "10.6%"), ("RMSE", "2.417", "2.557", "1.058"), ("Ljung Box p (lag-24)", "0.148", "-", "-")],
        "refs": [
            "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
            "NASA Langley Research Center. (n.d.). NASA POWER: Prediction of worldwide energy resources. https://power.larc.nasa.gov/",
            "Wickham, H., Francois, R., Henry, L., & Muller, K. (2023). dplyr: A grammar of data manipulation [Computer software]. https://CRAN.R-project.org/package=dplyr",
            "United Nations. (n.d.). Sustainable Development Goal 13: Climate action. https://sdgs.un.org/goals/goal13"
        ]
    }
}

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    r.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_text(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    r.font.size = Pt(12)
    return p

def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    r.font.size = Pt(12)
    return p

def add_figure(doc, image_path, caption):
    add_text(doc, caption)
    if image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(6.2))
    else:
        add_text(doc, "Figure file unavailable at report-generation time.")

def apply_default_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(12)
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

def make_report(key, cfg):
    doc = Document()
    apply_default_font(doc)

    # Cover page mirrors the reference report's centred institutional layout.
    for text, size, bold in [
        ("FACULTY OF", 14, True), ("KUALA LUMPUR CAMPUS", 14, True),
        ("DEPARTMENT OF MATHEMATICAL SCIENCES", 13, True),
        ("BMMS2094 Statistics for Data Science", 13, True),
        ("Assignment Individual", 13, True),
        ("Title Forecasting Monthly Mean Precipitation in Kuala Lumpur", 12, True)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(9)
        run = p.add_run(text)
        run.bold = bold; run.font.name = 'Times New Roman'; run.font.size = Pt(size)
    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    cover = [
        ("Student Name", cfg['student']), ("Student ID", "[To be completed]"),
        ("Academic Session", "202605"), ("Programme", "Bachelor in Data Science"),
        ("Tutorial Group", "6"), ("Lecturer's Name", "Assoc. Prof. Dr Chin Wan Yoke"),
        ("Practical Tutor's Name", "Assoc. Prof. Dr Chin Wan Yoke"),
        ("Date of submission", "05 September 2026")]
    for row, (left, right) in zip(table.rows, cover):
        set_cell_text(row.cells[0], left, True); set_cell_text(row.cells[1], right)
        set_cell_shading(row.cells[0], 'E7E6E6')
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Declaration: "); r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(12)
    add_text(doc, "I declare that this submitted coursework is my own work and that all material from other sources has been acknowledged appropriately.")
    add_text(doc, "Signature: ______________________________")
    doc.add_page_break()

    add_heading(doc, "I. METHODOLOGY")
    add_heading(doc, "A. Dataset Characteristics Motivating Model Selection")
    add_text(doc, cfg['method'][0], indent=True)
    add_heading(doc, "B. Model Specification and Parameter Selection")
    add_text(doc, cfg['method'][1], indent=True)
    add_heading(doc, "II. DATA ANALYSIS")
    add_text(doc, cfg['analysis'][0], indent=True)
    add_text(doc, cfg['analysis'][1], indent=True)
    add_heading(doc, "III. CONCLUSION")
    add_text(doc, cfg['conclusion'][0], indent=True)
    add_text(doc, cfg['conclusion'][1], indent=True)

    doc.add_page_break()
    add_heading(doc, "REFERENCES")
    for ref in cfg['refs']:
        add_text(doc, ref, indent=True)

    doc.add_page_break()
    add_heading(doc, "APPENDIX A: TABLE AND FIGURES")
    add_text(doc, f"Table 1. Final {cfg['model']} model accuracy, overfitting measures and residual whiteness.")
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric", "Training", "Cross validation", "MASE gap % / RMSE Ratio"]
    for c, h in zip(tbl.rows[0].cells, headers):
        set_cell_text(c, h, True); set_cell_shading(c, 'D9E2F3')
    for values in cfg['metrics']:
        cells = tbl.add_row().cells
        for c, value in zip(cells, values): set_cell_text(c, value)
    add_figure(doc, cfg['fig1'], cfg['fig1_caption'])
    add_figure(doc, cfg['fig2'], cfg['fig2_caption'])
    add_figure(doc, cfg['fig3'], cfg['fig3_caption'])

    doc.add_page_break()
    add_heading(doc, f"APPENDIX B: PROGRAMMING SOURCE CODE ({cfg['script']})")
    code_path = ROOT / 'scripts' / cfg['script']
    code = code_path.read_text(encoding='utf-8') if code_path.exists() else '# Source file unavailable.'
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        r.font.name = 'Courier New'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New'); r.font.size = Pt(8)

    doc.save(OUT / cfg['file'])

for key, cfg in MODELS.items():
    make_report(key, cfg)
    print(OUT / cfg['file'])
